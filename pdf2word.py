from io import BytesIO
from docx import Document
from docx.shared import Pt
import pdfplumber
import pytesseract
from pdf2image import convert_from_bytes
import camelot


def pdf_to_word(file, use_ocr=False):
    doc = Document()
    file.seek(0)

    if use_ocr:
        images = convert_from_bytes(file.read())
        for i, image in enumerate(images):
            text = pytesseract.image_to_string(image)
            for line in text.split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())
            if i < len(images) - 1:
                doc.add_page_break()
    else:
        with pdfplumber.open(file) as pdf:
            for pno, page in enumerate(pdf.pages):
                words = page.extract_words(extra_attrs=["size"])
                if not words:
                    continue

                avg_size = sum(w["size"] for w in words) / len(words)
                para = ""

                for w in words:
                    text = w["text"]
                    size = w["size"]

                    if size > avg_size + 2:
                        if para.strip():
                            doc.add_paragraph(para.strip())
                            para = ""
                        h = doc.add_heading(text, level=1)
                        h.style.font.size = Pt(size)
                    else:
                        para += text + " "

                if para.strip():
                    doc.add_paragraph(para.strip())

                if pno < len(pdf.pages) - 1:
                    doc.add_page_break()

        # table extraction (best-effort)
        try:
            tables = camelot.read_pdf(file, pages="all")
            for table in tables:
                df = table.df
                t = doc.add_table(rows=df.shape[0], cols=df.shape[1])
                for i in range(df.shape[0]):
                    for j in range(df.shape[1]):
                        t.rows[i].cells[j].text = df.iloc[i, j]
        except:
            pass

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
